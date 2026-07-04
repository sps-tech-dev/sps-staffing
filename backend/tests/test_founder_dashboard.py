"""B.11 founder dashboard: the role gate, tenant scoping, aggregate correctness,
cache behavior + PII-free structural guard (B4 invariant), by-bu zeros, audit rows,
exports."""
from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest
from argon2 import PasswordHasher
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from app import reporting
from app.db import get_redis, get_sessionmaker
from app.main import app
from app.models import (
    AuditLog, BusinessUnit, ClientUser, Consent, Membership, Notification, Tenant, User,
)
from app.models_staffing import (
    Application, Candidate, CandidateTimeline, Client, Invoice, Job, Offer, Placement, Test,
)

HOST = {"host": "spstechnosoft.com"}
PW = "FounderLocal!123"
FOUNDER = "fdr-owner@local.test"
RECRUITER = "fdr-rec@local.test"


def _mk_user(db, tenant, email, roles):
    old = db.execute(select(User).where(User.tenant_id == tenant.id,
                                        User.email == email)).scalar_one_or_none()
    if old is not None:
        db.execute(delete(Membership).where(Membership.user_id == old.id))
        db.execute(delete(ClientUser).where(ClientUser.user_id == old.id))
        db.execute(delete(User).where(User.id == old.id))
    u = User(tenant_id=tenant.id, email=email, password_hash=PasswordHasher().hash(PW),
             full_name="Fdr Tester", status="active")
    db.add(u); db.flush()
    bu = db.execute(select(BusinessUnit).where(BusinessUnit.tenant_id == tenant.id,
                                               BusinessUnit.code == "STAFFING")).scalar_one()
    db.add(Membership(user_id=u.id, business_unit_id=bu.id, roles=roles))
    db.commit()
    return u


def _flush_dash_cache(tenant_id):
    r = get_redis()
    for k in r.scan_iter(f"dash:founder:{tenant_id}:*"):
        r.delete(k)


@pytest.fixture
def env():
    db = get_sessionmaker()()
    sps = db.execute(select(Tenant).where(Tenant.code == "SPS001")).scalar_one()
    fdr = _mk_user(db, sps, FOUNDER, ["owner"])
    rec = _mk_user(db, sps, RECRUITER, ["recruiter"])
    _flush_dash_cache(sps.id)
    yield sps, db
    _flush_dash_cache(sps.id)
    db.execute(delete(Notification).where(Notification.tenant_id == sps.id))
    db.execute(delete(Test).where(Test.tenant_id == sps.id))
    db.execute(delete(Invoice).where(Invoice.tenant_id == sps.id))
    db.execute(delete(Placement).where(Placement.tenant_id == sps.id))
    db.execute(delete(CandidateTimeline).where(CandidateTimeline.tenant_id == sps.id))
    db.execute(delete(Consent).where(Consent.tenant_id == sps.id))
    for M in (Offer, Application, Job, Candidate, Client):
        db.execute(delete(M).where(M.tenant_id == sps.id))
    for u in (fdr, rec):
        db.execute(delete(Membership).where(Membership.user_id == u.id))
        db.execute(delete(User).where(User.id == u.id))
    db.commit(); db.close()


def _login(c, email, host=HOST):
    assert c.post("/api/auth/login", json={"email": email, "password": PW},
                  headers=host).status_code == 200


def _seed_commercials(db, sps, n_placements=2, ctc=1_000_000):
    client = Client(tenant_id=sps.id, business_unit_id="STAFFING", name="Fdr SeedCo")
    db.add(client); db.flush()
    job = Job(tenant_id=sps.id, business_unit_id="STAFFING", title="Fdr Seed Job",
              client_id=client.id)
    db.add(job); db.flush()
    today = dt.datetime.now(dt.timezone.utc).date()
    for i in range(n_placements):
        cand = Candidate(tenant_id=sps.id, full_name=f"Fdr Seed Cand {i}")
        db.add(cand); db.flush()
        appn = Application(tenant_id=sps.id, business_unit_id="STAFFING", job_id=job.id,
                           candidate_id=cand.id, client_id=client.id, stage="joined")
        db.add(appn); db.flush()
        plc = Placement(tenant_id=sps.id, business_unit_id="STAFFING",
                        application_id=appn.id, client_id=client.id, candidate_id=cand.id,
                        offered_ctc=ctc, joined_on=today,
                        guarantee_until=today + dt.timedelta(days=60))
        db.add(plc); db.flush()
        db.add(Invoice(tenant_id=sps.id, business_unit_id="STAFFING",
                       application_id=appn.id, client_id=client.id, placement_id=plc.id,
                       base_amount=ctc, fee_percent=15, fee_amount=ctc * 0.15,
                       total_amount=ctc * 0.15, status="draft"))
    db.commit()


def test_role_gate(env):
    sps, db = env
    _cid = None
    # founder (owner slug) → 200
    f = TestClient(app); _login(f, FOUNDER)
    assert f.get("/api/dashboard/founder/overview", headers=HOST).status_code == 200
    # recruiter → 403
    r = TestClient(app); _login(r, RECRUITER)
    assert r.get("/api/dashboard/founder/overview", headers=HOST).status_code == 403
    # plain admin → 403 (stricter than the admin gate, by design)
    adm = _mk_user(db, sps, "fdr-admin@local.test", ["admin"])
    try:
        a = TestClient(app); _login(a, "fdr-admin@local.test")
        assert a.get("/api/dashboard/founder/overview", headers=HOST).status_code == 403
    finally:
        db.execute(delete(Membership).where(Membership.user_id == adm.id))
        db.execute(delete(User).where(User.id == adm.id)); db.commit()
    # client session → 403
    client = Client(tenant_id=sps.id, business_unit_id="STAFFING", name="Fdr ClientCo")
    db.add(client); db.flush()
    u = User(tenant_id=sps.id, email="fdr-client@local.test",
             password_hash=PasswordHasher().hash(PW), full_name="FC", status="active")
    db.add(u); db.flush()
    db.add(ClientUser(tenant_id=sps.id, user_id=u.id, client_id=client.id, status="active"))
    db.commit()
    try:
        cc = TestClient(app); _login(cc, "fdr-client@local.test")
        assert cc.get("/api/dashboard/founder/overview", headers=HOST).status_code == 403
    finally:
        db.execute(delete(ClientUser).where(ClientUser.user_id == u.id))
        db.execute(delete(User).where(User.id == u.id))
        db.execute(delete(Client).where(Client.id == client.id)); db.commit()
    # unauth → 401
    assert TestClient(app).get("/api/dashboard/founder/overview",
                               headers=HOST).status_code == 401


def test_aggregates_match_hand_computed(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=2, ctc=1_000_000)
    f = TestClient(app); _login(f, FOUNDER)
    data = f.get("/api/dashboard/founder/overview", headers=HOST,
                 params={"refresh": 1}).json()
    assert data["revenue_mtd"] == 300000.0            # 2 × (1,000,000 × 15%)
    assert data["fees_billed_total"] == 300000.0
    assert data["placements_total"] == 2
    assert data["placements_in_guarantee"] == 2       # joined today, window open
    assert data["pipeline_funnel"]["joined"] == 2
    assert data["candidates_total"] == 2
    assert data["open_jobs"] == 1


def test_tenant_scoping(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=1)
    tb = db.execute(select(Tenant).where(Tenant.code == "TESTB13")).scalar_one_or_none()
    if tb is None:
        tb = Tenant(code="TESTB13", slug="testb13", name="Tenant B13"); db.add(tb); db.flush()
        db.add(BusinessUnit(tenant_id=tb.id, code="STAFFING", name="Staffing")); db.flush()
        db.commit()
    ub = _mk_user(db, tb, "fdr-b@local.test", ["owner"])
    B_HOST = {"host": "testb13.spstechnosoft.com"}
    try:
        _flush_dash_cache(tb.id)
        cb = TestClient(app); _login(cb, "fdr-b@local.test", B_HOST)
        data = cb.get("/api/dashboard/founder/overview", headers=B_HOST,
                      params={"refresh": 1}).json()
        assert data["placements_total"] == 0 and data["fees_billed_total"] == 0.0, \
            "LEAK: tenant B founder sees tenant A numbers"
    finally:
        db.execute(delete(Membership).where(Membership.user_id == ub.id))
        db.execute(delete(User).where(User.id == ub.id))
        db.execute(delete(BusinessUnit).where(BusinessUnit.tenant_id == tb.id))
        db.execute(delete(Tenant).where(Tenant.id == tb.id)); db.commit()


def test_cache_hit_and_refresh(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=1)
    f = TestClient(app); _login(f, FOUNDER)
    first = f.get("/api/dashboard/founder/overview", headers=HOST,
                  params={"refresh": 1}).json()
    assert first["placements_total"] == 1
    _seed_commercials(db, sps, n_placements=1)        # mutate underlying data
    cached = f.get("/api/dashboard/founder/overview", headers=HOST).json()
    assert cached["placements_total"] == 1, "cache did not serve within TTL"
    fresh = f.get("/api/dashboard/founder/overview", headers=HOST,
                  params={"refresh": 1}).json()
    assert fresh["placements_total"] == 2, "?refresh=1 did not recompute"


def test_pii_free_cache_guard(env):
    # structural B4 proof: names/emails can NEVER reach a founder cache value
    with pytest.raises(reporting.CacheGuardError):
        reporting.cache_guard({"top_recruiter": "Asha Sharma"})
    with pytest.raises(reporting.CacheGuardError):
        reporting.cache_guard({"contact": "a@b.example"})
    with pytest.raises(reporting.CacheGuardError):
        reporting.cache_guard(["ok", {"note": "free text sentence here"}])
    # real payload shapes pass: enumerated field keys, label/date data keys, numeric leaves
    reporting.cache_guard({"revenue_mtd": 1.5, "pipeline_funnel": {"applied": 3},
                           "series": {"2026-07": 2}, "metric": "revenue",
                           "assessment_pass_rate": None,
                           "notifications": {"sent": 1, "pending": 0}})


def test_by_bu_zeros_not_fabricated(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=1)
    f = TestClient(app); _login(f, FOUNDER)
    data = f.get("/api/dashboard/founder/by-bu", headers=HOST, params={"refresh": 1}).json()
    assert data["STAFFING"]["placements"] == 1
    for bu in ("ACADEMY", "CONSULTING"):
        assert data[bu] == {"jobs": 0, "applications": 0, "placements": 0, "fees_billed": 0.0}


def test_trends_and_audit_rows(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=1)
    before = db.execute(select(func.count()).select_from(AuditLog).where(
        AuditLog.tenant_id == sps.id,
        AuditLog.action == "dashboard.founder_access")).scalar_one()
    f = TestClient(app); _login(f, FOUNDER)
    t = f.get("/api/dashboard/founder/trends", headers=HOST,
              params={"metric": "placements", "range": 3, "refresh": 1}).json()
    assert t["metric"] == "placements" and len(t["series"]) == 3
    assert sum(t["series"].values()) == 1
    assert f.get("/api/dashboard/founder/trends", headers=HOST,
                 params={"metric": "nope"}).status_code == 422
    f.get("/api/dashboard/founder/overview", headers=HOST)
    after = db.execute(select(func.count()).select_from(AuditLog).where(
        AuditLog.tenant_id == sps.id,
        AuditLog.action == "dashboard.founder_access")).scalar_one()
    # exactly TWO audited accesses: the trends call + the overview call;
    # the 422 bad-metric request never reaches _audited
    assert after - before == 2


def test_exports(env):
    sps, db = env
    _seed_commercials(db, sps, n_placements=1)
    f = TestClient(app); _login(f, FOUNDER)
    pdf = f.get("/api/dashboard/founder/export", headers=HOST, params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    from pdfminer.high_level import extract_text
    text = extract_text(io.BytesIO(pdf.content))
    assert "revenue_mtd" in text and "Fdr Seed Cand" not in text     # aggregates only
    xlsx = f.get("/api/dashboard/founder/export", headers=HOST, params={"format": "xlsx"})
    assert xlsx.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert set(wb.sheetnames) == {"Overview", "By BU"}
    flat = [str(c.value) for row in wb["Overview"].iter_rows() for c in row if c.value]
    assert not any("Fdr Seed Cand" in v for v in flat)               # no names anywhere
    # recruiter cannot export
    r = TestClient(app); _login(r, RECRUITER)
    assert r.get("/api/dashboard/founder/export", headers=HOST).status_code == 403
