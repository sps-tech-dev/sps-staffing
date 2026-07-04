"""Resume text extraction (B.1) — pure functions, no AWS, no AI.

PDF  → pdfminer.six; DOCX → python-docx, docx2txt fallback; legacy .doc → best-effort
empty (no reliable pure-Python extractor for the OLE .doc format — upload still
succeeds, text is just absent). All pure-Python deps (ARM64-safe slim base).

Output is capped and sanitized: control characters stripped (keeps \\n and \\t),
length bounded so a pathological file can't bloat the row. Extraction failures
return "" — the upload is never failed over unparseable content.
"""
from __future__ import annotations

import io
import logging
import re

log = logging.getLogger("sps.resume_parse")

MAX_TEXT_CHARS = 200_000
# strip C0/C1 control chars except \n (0A) and \t (09); \r normalized first
_CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f]")

PDF = "application/pdf"
DOC = "application/msword"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _sanitize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    # collapse runs of blank lines left behind by layout extraction
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:MAX_TEXT_CHARS]


def _extract_pdf(data: bytes) -> str:
    from pdfminer.high_level import extract_text
    return extract_text(io.BytesIO(data)) or ""


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        if any(p.strip() for p in parts):
            return "\n".join(parts)
    except Exception:  # noqa: BLE001 — fall through to docx2txt
        pass
    import docx2txt
    return docx2txt.process(io.BytesIO(data)) or ""


def extract_text(data: bytes, content_type: str) -> str:
    """Extract plain text from a resume. Returns "" when the format has no extractor
    (.doc) or the file is unparseable — callers treat text as best-effort."""
    try:
        if content_type == PDF:
            return _sanitize(_extract_pdf(data))
        if content_type == DOCX:
            return _sanitize(_extract_docx(data))
        if content_type == DOC:
            log.info("legacy .doc upload — no pure-python extractor; storing without text")
            return ""
        raise ValueError(f"unsupported resume content-type: {content_type}")
    except ValueError:
        raise
    except Exception as e:  # noqa: BLE001 — never fail the upload on a bad parse
        log.warning("resume text extraction failed (%s): %s", content_type, e)
        return ""
